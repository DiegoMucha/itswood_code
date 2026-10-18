"""Genera hoja/pestanas_bot_itswood.xlsx: una pestaña por necesidad lógica, con las cabeceras de columnas.json
y BOT_CATALOGO desde catalogo_bot.csv. Las pestañas se importan al libro de ITS WOOD con
Archivo > Importar > Subir > "Insertar hojas nuevas", sin tocar las pestañas existentes.

Uso: python hoja/generar_plantilla.py   (después de node n8n/build.js, que actualiza columnas.json)
"""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.comments import Comment

DIR = Path(__file__).parent
TEXTO = "@"
COLUMNAS_TEXTO = {"wa_id", "telefono", "mensaje_id", "id", "logo_media_id", "media_id", "fecha_deseada", "valor"}
COLUMNAS_EQUIPO = {"responsable", "notas", "notas_equipo"}
NOTAS = {
    "modo": "BOT, PENDIENTE_HUMANO, HUMANO o CERRADO. Cámbialo con el formulario de atención, no a mano.",
    "responsable": "Columna del equipo: el bot no la escribe.",
    "notas": "Columna del equipo: el bot no la escribe.",
    "notas_equipo": "Columna del equipo: el bot no la escribe.",
    "estado": "El bot la crea en PENDIENTE_HUMANO; después la gestiona el equipo.",
    "fecha_confirmada": "PENDIENTE hasta que el equipo confirme la fecha de producción.",
    "minimo": "Opcional: mínimo de unidades autorizado para el producto. Vacío = sin mínimo.",
    "activo": "Opcional: escribe NO para ocultar el producto al bot.",
    "nombre_interno": "Nombre exacto en NOMBRES / INVENTARIO. De aquí salen el stock y el producto de PAUTAS.",
    "stock": "Solo se usa si el producto no figura en INVENTARIO.",
}


def cabecera(ws, columnas):
    ws.append(columnas)
    for i, nombre in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="7A5C3E" if nombre not in COLUMNAS_EQUIPO else "3E6B7A")
        if nombre in NOTAS:
            celda.comment = Comment(NOTAS[nombre], "ITS WOOD bot")
        ws.column_dimensions[celda.column_letter].width = max(12, min(40, len(nombre) + 6))
        if nombre in COLUMNAS_TEXTO:
            # Texto: evita que los números de teléfono se conviertan en notación científica.
            for fila in range(2, 2001):
                ws.cell(row=fila, column=i).number_format = TEXTO
    ws.freeze_panes = "A2"


def main():
    columnas = json.loads((DIR / "columnas.json").read_text(encoding="utf-8"))
    iniciales = json.loads((DIR / "filas_iniciales.json").read_text(encoding="utf-8"))
    wb = Workbook()
    wb.remove(wb.active)
    for pestana, cols in columnas.items():
        ws = wb.create_sheet(pestana)
        cabecera(ws, cols)
        for fila in iniciales.get(pestana, []):
            ws.append([fila.get(c, "") for c in cols])

    with open(DIR / "catalogo_bot.csv", encoding="utf-8") as f:
        filas = list(csv.reader(f))
    ws = wb.create_sheet("BOT_CATALOGO")
    cabecera(ws, filas[0])
    for fila in filas[1:]:
        if not fila:
            continue
        precio = float(fila[2]) if fila[2] else None
        stock = int(fila[3]) if fila[3] else None
        minimo = int(fila[5]) if fila[5] else None
        ws.append([fila[0], fila[1], precio, stock, fila[4] or None, minimo, fila[6] or None, fila[7] or None])
    ws.column_dimensions["B"].width = 55
    ws.column_dimensions["E"].width = 30
    ws.column_dimensions["H"].width = 60

    destino = DIR / "pestanas_bot_itswood.xlsx"
    wb.save(destino)
    print(f"{destino.name}: {', '.join(wb.sheetnames)}; {len(filas) - 1} productos")


if __name__ == "__main__":
    main()
